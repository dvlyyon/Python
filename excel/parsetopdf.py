#!/usr/bin/python3

import sys
import re
from enum import IntEnum
from openpyxl import Workbook
import datetime

from openpyxl.chart import (LineChart, Reference,)
from openpyxl.chart.axis import DateAxis

class State(IntEnum):
    INIT = 1
    DATE = 2
    T_CPU = 3
    T_P_MEM = 4
    T_L_MEM = 5
    T_CPU_MEM = 6
    D_STORAGE = 7


def appendData(toList,fromDic, keys):
    for name in keys:
        toList.append(fromDic[name])

def appendTitle(toTitle, suffix, keys):
    for name in keys:
        toTitle.append(name[0:-4] + suffix)

def cast2K(value):
    if value[-1] == 'm':
        tmpValue=float(value[0:-1])
        tmpValue = tmpValue * 1024
    elif value[-1] == 'g':
        tmpValue=float(value[0:-1])
        tmpValue = tmpValue * 1024 * 1024
    else:
        tmpValue=float(value)
    return tmpValue

colorsPatten=[ "FF0000", "AA0000", "00FF00", "00AA00", "0000FF",  "9000FF", "00EEEE","B58215",
               "FFAA00", "E516EB", "535BAD", "AA00FF", "00AAFF", "FFFF00", "AA00AA", "00AAAA" ]

def drawLineChartInChartSheet(worksheet, chartSheet, data, minCol, maxCol, title, yTitle, xTitle, position):
    chartData = Reference(worksheet,min_col=minCol,min_row=1,max_col=maxCol,max_row=len(data))

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title=yTitle
    chart.y_axis.majorGridlines=None
    chart.x_axis.title=xTitle

    chart.add_data(chartData,titles_from_data=True)

    for n in range(0,maxCol-minCol+1):
        style = chart.series[n]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 22000
        style.smooth=True
    chartSheet.add_chart(chart) 

def drawLineChart(worksheet, data, minCol, maxCol, title, yTitle, xTitle, position):
    chartData = Reference(worksheet,min_col=minCol,min_row=1,max_col=maxCol,max_row=len(data))

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title=yTitle
    chart.y_axis.majorGridlines=None
    chart.x_axis.title=xTitle

    chart.add_data(chartData,titles_from_data=True)

    for n in range(0,maxCol-minCol+1):
        style = chart.series[n]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 22000
        style.smooth=True
    worksheet.add_chart(chart, "A"+str(len(data)+position))

origFile = open(sys.argv[1])

wb = Workbook()
ws = wb.active
cs = wb.create_chartsheet(title='CPU_MEM_DISK Chart',index=0)

ws.title = "CPU_MEM_DISK Data"

title = ["Time", "Total CPU Used%", "Total (P)Memory Used%", "Total (L)Memory Used%","Total Free Mem(KB)","Total Avail Mem(KB)"]
data=[]
data.append(title)
state=State.INIT
tmpline=[0,0,0,0,0,0] 
totalmem=0
processNum=0
processNameList=[]
processDataVIR={}
processDataRES={}
processDataCPU={}
processDataMEM={}
processCollected=False
entry=State.INIT
for line in origFile: 
    if state == State.INIT:
        match = re.match("^date$", line)
        if match: 
            state=State.DATE 
    elif state == State.DATE:
        timefields = re.split('[ ]+',line.strip())
        datestr = timefields[1]+' '+timefields[2]+' '+ \
                  timefields[3]+' '+timefields[5]
        print(datestr)
        date = datetime.datetime.strptime(datestr,'%b %d %H:%M:%S %Y')
        tmpline[0]=date
        state = State.T_CPU
    elif state == State.T_CPU:
        match = re.match('^%Cpu\(s\):\s+.*\s+([0-9.]+)\s+id,.*st$',line)
        if match:
            print('Match Total CPU')
            tmpline[1]=round(float(100-float(match.group(1))),1)
            state=State.T_P_MEM
    elif state == State.T_P_MEM:
        match = re.match('^KiB\sMem\s+:\s+(\d+)\s+total,\s+(\d+)\s+free,\s+(\d+)\s+used,\s+(\d+)\s+buff.cache$',line)
        if match:
            print('Match Total P Mem')
            totalmem = int(match.group(1))
            tmpline[2]=round(float((int(match.group(1))-int(match.group(2)))/int(match.group(1))*100),1)
            tmpline[4]=int(match.group(2))
            state=State.T_L_MEM
    elif state == State.T_L_MEM:
        match = re.match('^KiB\sSwap:\s+.*used\.\s+(\d+)\s+avail\s+Mem\s*$',line)
        if match:
            print('Match Total L Mem')
            tmpline[3]=round(float((totalmem-int(match.group(1)))/totalmem*100),1)
            tmpline[5]=int(match.group(1))
            state=State.T_CPU_MEM
            entry=State.T_L_MEM
    elif state == State.T_CPU_MEM:
        match = re.match('^\s*(\d+)\s+root\s+(\d+)\s+(\d+)\s+([\d.]+[mg]?)\s+([\d.]+[mg]?)\s+(\d+)\s+\w+\s+([0-9.]+)\s+([0-9.]+)\s+[0-9:.]+\s+(\w+\.bin)',line)
        if match: 
            processName=match.group(9)
            print('Match ' + processName)
            if not processCollected:
                processNameList.append(processName)
            else:
                assert processName in processNameList
            virM=match.group(4)
            resM=match.group(5)
            processDataVIR[processName]=cast2K(virM)
            processDataRES[processName]=cast2K(resM)
            processDataCPU[processName]=float(match.group(7))
            processDataMEM[processName]=float(match.group(8))
            entry=State.T_CPU_MEM
        else:
            if entry == State.T_CPU_MEM:
                state=State.D_STORAGE
                if not processCollected:
                    appendTitle(title,' CPU%',processNameList)
                    appendTitle(title,' MEM%',processNameList)
                    appendTitle(title,' VIR(KB)',processNameList)
                    appendTitle(title,' RES(KB)',processNameList)
                    title.append("DB Usage%")
                processCollected=True
                appendData(tmpline,processDataCPU,processNameList)
                appendData(tmpline,processDataMEM,processNameList)
                appendData(tmpline,processDataVIR,processNameList)
                appendData(tmpline,processDataRES,processNameList)
                processDataCPU.clear()
                processDataMEM.clear()
                processDataRES.clear()
                processDataVIR.clear()
    elif state == State.D_STORAGE:
        match = re.match('^.*\s+([0-9.]+)%\s+\/storage$',line)
        if match:
            print('Match storage')
            tmpline.append(int(match.group(1)))
            data.append(tmpline)
            print(tmpline)
            tmpline=[0,0,0,0,0,0]
            state = State.INIT
assert(state==State.INIT)
            
for row in data:
    ws.append(row)

processNum=len(processNameList)
#drawLineChart(cs,data,2,4,"Total CPU & Memory Usage","Usage %","Time No.", 5)
#drawLineChart(cs,data,5,6,"Total Free & Avail Memory", "Memory (KB)","Time No.",25)
#drawLineChart(cs,data,7,6+processNum,"CPU Usage per Process", "Usage %","Time No.", 45)
#drawLineChart(cs,data,7+processNum,6+processNum*2,"Memory Usage per Process", "Usage %","Time No.",65)
#drawLineChart(cs,data,7+processNum*2,6+processNum*3,"Virtual Memory per Process", "Memory (KB)","Time No.",85)
#drawLineChart(cs,data,7+processNum*3,6+processNum*4,"Resident Memory per Process", "Memory (KB)","Time No.",105)
#drawLineChart(cs,data,7+processNum*4,7+processNum*4,"Disk Usage", "Usage %","Time No.",125)
drawLineChartInChartSheet(ws,cs,data,2,4,"Total CPU & Memory Usage","Usage %","Time No.", 5)
drawLineChart(ws,data,5,6,"Total Free & Avail Memory", "Memory (KB)","Time No.",5)
drawLineChart(ws,data,7,6+processNum,"CPU Usage per Process", "Usage %","Time No.", 25)
drawLineChart(ws,data,7+processNum,6+processNum*2,"Memory Usage per Process", "Usage %","Time No.",45)
drawLineChart(ws,data,7+processNum*2,6+processNum*3,"Virtual Memory per Process", "Memory (KB)","Time No.",65)
drawLineChart(ws,data,7+processNum*3,6+processNum*4,"Resident Memory per Process", "Memory (KB)","Time No.",85)
drawLineChart(ws,data,7+processNum*4,7+processNum*4,"Disk Usage", "Usage %","Time No.",105)

wb.save(sys.argv[1]+".xlsx")
origFile.close()
