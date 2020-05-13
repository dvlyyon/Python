#!/usr/bin/python3

import sys
import re
from enum import IntEnum
from openpyxl import Workbook
import datetime

from openpyxl.chart import (LineChart, Reference,)
from openpyxl.chart.axis import DateAxis

class State(IntEnum):
    INIT = 1
    DATE = 2
    T_CPU = 3
    T_MEM = 4
    T_CPU_MEM = 5
    D_STORAGE = 6

origFile = open(sys.argv[1])

wb = Workbook()
ws = wb.active
ws.title = "Memory Statistics"

title = ["Time", "Total CPU Used%", "Total Memory Used%", "MF Memory %",  "SMF Memory %", "DBF Memory %", "MF CPU %", "SMF CPU %", "DBF CPU %", "Storage Used%"]
data=[]
data.append(title)
state=State.INIT
tmpline=[0,0,0,0,0,0,0,0,0,0] 
topmark='' 
for line in origFile: 
    if state == State.INIT:
        match = re.match("^date$", line)
        if match: 
            state=State.DATE 
    elif state == State.DATE:
        timefields = re.split('[ ]+',line.strip())
        datestr = timefields[1]+' '+timefields[2]+' '+ \
                  timefields[3]+' '+timefields[5]
        print(datestr)
        date = datetime.datetime.strptime(datestr,'%b %d %H:%M:%S %Y')
        tmpline[0]=date
        state = State.T_CPU
    elif state == State.T_CPU:
        match = re.match('^%Cpu\(s\):\s+.*\s+([0-9.]+)\s+id,.*st$',line)
        if match:
            print('Match Total CPU')
            tmpline[1]=round(float(100-float(match.group(1))),1)
            state=State.T_MEM
    elif state == State.T_MEM:
        match = re.match('^KiB\sMem\s+:\s+(\d+)\s+total,\s+(\d+)\s+free,.*cache$',line)
        if match:
            print('Match Total Mem')
            tmpline[2]=round(float((int(match.group(1))-int(match.group(2)))/int(match.group(1))*100),1)
            state=State.T_CPU_MEM
    elif state == State.T_CPU_MEM:
        match1 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+MF.bin',line)
        match2 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+SMF.bin',line)
        match3 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+DBF.bin',line)
        if match1 or match2 or match3:
            if match1:
                print('Match Top MF')
                topmark=topmark+"X"
                tmpline[3]=float(match1.group(2))
                tmpline[6]=float(match1.group(1))
            elif match2:
                print('Match Top SMF')
                topmark=topmark+"S"
                tmpline[4]=float(match2.group(2))
                tmpline[7]=float(match2.group(1))
            elif match3:
                print('Match Top DBF')
                topmark=topmark+"Z"
                tmpline[5]=float(match3.group(2))
                tmpline[8]=float(match3.group(1))
            if topmark.find('X')>=0 and topmark.find('S')>=0 and topmark.find('Z')>=0:
                state=State.D_STORAGE
                topmark=''
    elif state == State.D_STORAGE:
        match = re.match('^.*\s+([0-9.]+)%\s+\/storage$',line)
        if match:
            print('Match storage')
            tmpline[9] = int(match.group(1))
            data.append(tmpline)
            print(tmpline)
            tmpline=[0,0,0,0,0,0,0,0,0,0]
            state = State.INIT
assert(state==State.INIT)
            
for row in data:
    ws.append(row)

data1 = Reference(ws,min_col=2, min_row=1, max_col=3, max_row=len(data))

c1 = LineChart()
c1.title = "Total CPU & Memory Information"
c1.style = 13
c1.y_axis.majorGridlines = None
c1.y_axis.title = "CPU & Memory Usage %"
c1.x_axis.title = "Time"

c1.add_data(data1, titles_from_data=True)

s1 = c1.series[0]
s1.graphicalProperties.line.solidFill = "0000FF"
s1.graphicalProperties.line.width = 30000 # width in EMUs

s2 = c1.series[1]
s2.graphicalProperties.line.solidFill = "00FF00"
s2.graphicalProperties.line.width = 30000 # width in EMUs
s2.smooth = True
ws.add_chart(c1, "A"+str(len(data)+5))

data2 = Reference(ws,min_col=4, min_row=1, max_col=6, max_row=len(data))
c2 = LineChart()
c2.title = "Individule Memory Information"
c2.style = 13
c2.y_axis.majorGridlines = None
c2.y_axis.title = "Memory Usage %"
c2.x_axis.title = "Time"

c2.add_data(data2, titles_from_data=True)

s21 = c2.series[0]
s21.graphicalProperties.line.solidFill = "0000FF"
s21.graphicalProperties.line.width = 30000 # width in EMUs

s22 = c2.series[1]
s22.graphicalProperties.line.solidFill = "00FF00"
s22.graphicalProperties.line.width = 30000 # width in EMUs

s23 = c2.series[2]
s23.graphicalProperties.line.solidFill = "FF0000"
s23.graphicalProperties.line.width = 30000 # width in EMUs

ws.add_chart(c2, "A"+str(len(data)+25))


data3 = Reference(ws,min_col=7, min_row=1, max_col=9, max_row=len(data))
c3 = LineChart()
c3.title = "Individule CPU Information"
c3.style = 13
c3.y_axis.majorGridlines = None
c3.y_axis.title = "CPU Usage %"
c3.x_axis.title = "Time"

c3.add_data(data3, titles_from_data=True)

s31 = c3.series[0]
s31.graphicalProperties.line.solidFill = "0000FF"
s31.graphicalProperties.line.width = 30000 # width in EMUs

s32 = c3.series[1]
s32.graphicalProperties.line.solidFill = "00FF00"
s32.graphicalProperties.line.width = 30000 # width in EMUs

s33 = c3.series[2]
s33.graphicalProperties.line.solidFill = "FF0000"
s33.graphicalProperties.line.width = 30000 # width in EMUs
ws.add_chart(c3, "A"+str(len(data)+45))

data4 = Reference(ws,min_col=10, min_row=1, max_col=10, max_row=len(data))
c4 = LineChart()
c4.title = "Disk (storage) Usage"
c4.style = 13
c4.y_axis.majorGridlines = None
c4.y_axis.title = "Disk Usage"
c4.x_axis.title = "Time"
c4.add_data(data4, titles_from_data=True)

s41 = c4.series[0]
s41.graphicalProperties.line.solidFill = "0000FF"
s41.graphicalProperties.line.width = 30000 # width in EMUs

ws.add_chart(c4, "A"+str(len(data)+65))

wb.save(sys.argv[1]+".telemetry.xlsx")
origFile.close()
