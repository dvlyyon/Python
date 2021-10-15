#!/usr/bin/python3

import sys
import re
from enum import IntEnum
from openpyxl import Workbook
import datetime
import json

from openpyxl.chart import (LineChart, Reference,)
from openpyxl.chart.axis import DateAxis

class State(IntEnum):
    INIT = 1
    DATE = 2
    T_CPU = 3
    T_P_MEM = 4
    T_L_MEM = 5
    T_CPU_MEM = 6
    D_STORAGE = 7

class DateTimeEncoder(json.JSONEncoder):
    def default(self, o):
        if isinstance(o,datetime.datetime):
            return o.isoformat()
        return json.JSONEncoder.default(self,o)

def appendData(toList,fromDic, keys):
    for name in sorted(keys):
        toList.append(fromDic[name])

def appendProcessTitle(toTitle, suffix, keys):
    for name in sorted(keys):
        if name.endswith('.bin'):
            toTitle.append(name[0:-4] + suffix)
        else:
            toTitle.append(name + suffix)

def appendFileSystemTitle(toTitle, keys):
    for name in sorted(keys):
#        if name == '/':
#            tmpString='root'
#        else:
#            tmpString=name.replace('/',' ')
        toTitle.append('FS: '+name)

def cast2K(value):
    if value[-1] == 'm':
        tmpValue=float(value[0:-1])
        tmpValue = tmpValue * 1024
    elif value[-1] == 'g':
        tmpValue=float(value[0:-1])
        tmpValue = tmpValue * 1024 * 1024
    else:
        tmpValue=float(value)
    return tmpValue

colorsPatten=[ "FF0000", "00FF00", "0000FF", "00AA00",  "9000FF", "AA0000", "00EEEE", "B58215",
               "FFAA00", "E516EB", "535BAD", "AA00FF", "00AAFF", "FFFF00", "AA00AA", "00AAAA",
               "44EEAA", "37FF00", "FF2277", "00AA99"]

def drawLineChartInChartSheet(worksheet, chartSheet, data, minCol, maxCol, title, yTitle, xTitle, position):
    chartData = Reference(worksheet,min_col=minCol,min_row=1,max_col=maxCol,max_row=len(data))

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title=yTitle
    chart.y_axis.majorGridlines=None
    chart.x_axis.title=xTitle

    chart.add_data(chartData,titles_from_data=True)
    chart.legend.position='t'

    for n in range(0,maxCol-minCol+1):
        style = chart.series[n]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 22000
        style.smooth=True
    chartSheet.add_chart(chart) 

def drawMixLineChartInChartSheet(worksheet, chartSheet, data, minCol, maxCol, title, yTitle, xTitle, minCol1, maxCol1, yTitle1):
    chartData = Reference(worksheet,min_col=minCol,min_row=1,max_col=maxCol,max_row=len(data))
    chartData1 = Reference(worksheet,min_col=minCol1,min_row=1,max_col=maxCol1,max_row=len(data))

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title=yTitle
    chart.y_axis.majorGridlines=None
    chart.x_axis.title=xTitle

    chart.add_data(chartData,titles_from_data=True)
    chart.legend.position='t'

    for n in range(0,maxCol-minCol+1):
        style = chart.series[n]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 22000
        style.smooth=True

    chart1 = LineChart()
    chart1.add_data(chartData1,titles_from_data=True)
    chart1.y_axis.axId = 200
    chart1.y_axis.title=yTitle1
    chart1.x_axis.title=xTitle

    chart.legend.position='t'

    for n in range(maxCol-minCol+1,maxCol-minCol+2+maxCol1-minCol1):
        style = chart1.series[n-maxCol+minCol-1]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 40000
        style.smooth=True

    chart1.y_axis.crosses = "max"
    chart += chart1

    chartSheet.add_chart(chart) 

def drawLineChart(worksheet, data, minCol, maxCol, title, yTitle, xTitle, position):
    chartData = Reference(worksheet,min_col=minCol,min_row=1,max_col=maxCol,max_row=len(data))

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title=yTitle
    chart.y_axis.majorGridlines=None
    chart.x_axis.title=xTitle

    chart.add_data(chartData,titles_from_data=True)
    chart.legend.position='t'
    chart.width=30
    chart.height=12.5
    
    for n in range(0,maxCol-minCol+1):
        style = chart.series[n]
        if n < len(colorsPatten):
            style.graphicalProperties.line.solidFill = colorsPatten[n]
        style.graphicalProperties.line.width = 22000
        style.smooth=True
    worksheet.add_chart(chart, "A"+str(len(data)+position))

origFile = open(sys.argv[1])

wb = Workbook()
ws = wb.active
csTotal = wb.create_chartsheet(title='Total CPU & Memory',index=0)
csDisk = wb.create_chartsheet(title='Dist Usage', index=1)

ws.title = "Detailed CPU Memory"

title = ["Time", "Total CPU Used%", "Total (P)Memory Used%", "Total (L)Memory Used%","Total Free Mem(KB)","Total Avail Mem(KB)"]
data=[]
data.append(title)
state=State.INIT
tmpline=[0,0,0,0,0,0] 
totalmem=0
processNum=0
processNameList=[]
processDataVIR={}
processDataRES={}
processDataCPU={}
processDataMEM={}
processCollected=False
mountPointCollected=False
mountPointList=[]
mountPointData={}
entry=State.INIT
newDate=True
lineNum=0
for line in origFile: 
    lineNum += 1
    if state == State.INIT:
        match = re.match("^date -Iseconds$", line)
        if match:
            newDate = True
        else:
            match = re.match("^date$", line)
            if match:
                newDate = False
        if match: 
            state=State.DATE 
    elif state == State.DATE:
        timestr = line.strip()
        print(timestr)
        if newDate:
            datestr = timestr[:19]+timestr[19:].replace(":","")
            date = datetime.datetime.strptime(datestr,'%Y-%m-%dT%H:%M:%S%z')
        else:
            timefields = re.split('[ ]+',timestr)
            datestr = timefields[1]+' '+timefields[2]+' '+ \
                      timefields[3]+' '+timefields[5]

            date = datetime.datetime.strptime(datestr,'%b %d %H:%M:%S %Y')
        tmpline[0]=date
        state = State.T_CPU
    elif state == State.T_CPU:
        match = re.match('^%Cpu\(s\):\s+.*([0-9.]+)\s+id,.*st$',line)
        if match:
            print('Match Total CPU')
            tmpline[1]=round(float(100-float(match.group(1))),1)
            state=State.T_P_MEM
        else:
            match = re.match("^date -Iseconds$", line)
            if match:
                break
    elif state == State.T_P_MEM:
        match = re.match('^[KM]iB\sMem\s+:\s+([\d.]+)\s+total,\s+([\d.]+)\s+free,\s+([\d.]+)\s+used,\s+([\d.]+)\s+buff.cache$',line)
        if match:
            print('Match Total P Mem')
            totalmem = float(match.group(1))
            tmpline[2]=round(float((float(match.group(1))-float(match.group(2)))/float(match.group(1))*100),1)
            tmpline[4]=float(match.group(2))
            state=State.T_L_MEM
    elif state == State.T_L_MEM:
        match = re.match('^[KM]iB\sSwap:\s+.*used\.\s+([\d.]+)\s+avail\s+Mem\s*$',line)
        if match:
            print('Match Total L Mem')
            tmpline[3]=round(float((totalmem-float(match.group(1)))/totalmem*100),1)
            tmpline[5]=float(match.group(1))
            jsonline=tmpline.copy()
            state=State.T_CPU_MEM
            entry=State.T_L_MEM
    elif state == State.T_CPU_MEM:
        match = re.match('^\s*(\d+)\s+root\s+(\d+)\s+(\d+)\s+([\d.]+[mg]?)\s+([\d.]+[mg]?)\s+(\d+)\s+\w+\s+([0-9.]+)\s+([0-9.]+)\s+[0-9:.]+\s+([\w.-]+)',line)
        if match: 
            processName=match.group(9)
            print('Match ' + processName)
            if not processCollected:
                processNameList.append(processName)
            else:
                assert processName in processNameList
            virM=match.group(4)
            resM=match.group(5)
            processDataVIR[processName]=cast2K(virM)
            processDataRES[processName]=cast2K(resM)
            processDataCPU[processName]=float(match.group(7))
            processDataMEM[processName]=float(match.group(8))
            entry=State.T_CPU_MEM
        else:
            if entry == State.T_CPU_MEM:
                state=State.D_STORAGE
                if not processCollected:
                    appendProcessTitle(title,' CPU%',processNameList)
                    appendProcessTitle(title,' MEM%',processNameList)
                    appendProcessTitle(title,' VIR(KB)',processNameList)
                    appendProcessTitle(title,' RES(KB)',processNameList)
                processCollected=True
                appendData(tmpline,processDataCPU,processNameList)
                appendData(tmpline,processDataMEM,processNameList)
                appendData(tmpline,processDataVIR,processNameList)
                appendData(tmpline,processDataRES,processNameList)
                processDataCPU.clear()
                processDataMEM.clear()
                processDataRES.clear()
                processDataVIR.clear()
    elif state == State.D_STORAGE:
        match = re.match('^[\w\d/]+\s+\d+\s+\d+\s+\d+\s+([0-9.]+)%\s+(\/.*)$',line)
        if match:
            mountPoint=match.group(2)
            print('Match storage:' + mountPoint)
            if mountPointCollected:
                assert mountPoint in mountPointList
            else:
                mountPointList.append(mountPoint)
            mountPointData[mountPoint]=int(match.group(1))
            entry=State.D_STORAGE
        else:
            if entry == State.D_STORAGE:
                if not mountPointCollected:
                    appendFileSystemTitle(title,mountPointList)
                mountPointCollected=True
                appendData(tmpline,mountPointData,mountPointList)
                mountPointData.clear()
                data.append(tmpline)
                print(tmpline)
                assert len(tmpline) == len(title)
                tmpline=[0,0,0,0,0,0]
                state = State.INIT
print(state)
print(lineNum)
assert(state==State.INIT)

for row in data:
    ws.append(row)

with open(sys.argv[1]+".txt", 'w') as json_file:
    json.dump(data, json_file, cls=DateTimeEncoder)

processNum=len(processNameList)
mountPointNum=len(mountPointList)
drawMixLineChartInChartSheet(ws,csTotal,data,2,4,"Total CPU & Memory ","Usage %","Timestamp No.", 5, 6, "Size (KB)")
drawLineChartInChartSheet(ws,csDisk,data,7+processNum*4,6+processNum*4+mountPointNum,"Disk Usage", "Usage %", "Timestamp No.",5)
#drawLineChart(ws,data,2,4,"Total CPU & Memory Usage","Usage %","Timestamp No.", 5)
#drawLineChart(ws,data,5,6,"Total Free & Avail Memory", "Memory (KB)","Timestamp No.",5)
drawLineChart(ws,data,7,6+processNum,"CPU Usage per Process", "Usage %","Timestamp No.", 5)
drawLineChart(ws,data,7+processNum,6+processNum*2,"Memory Usage per Process", "Usage %","Timestamp No.",35)
drawLineChart(ws,data,7+processNum*2,6+processNum*3,"Virtual Memory per Process", "Memory (KB)","Timestamp No.",65)
drawLineChart(ws,data,7+processNum*3,6+processNum*4,"Resident Memory per Process", "Memory (KB)","Timestamp No.",95)
#drawLineChart(ws,data,7+processNum*4,6+processNum*4+mountPointNum,"Disk Usage", "Usage %","Timestamp No.",125)

wb.save(sys.argv[1]+".xlsx")
origFile.close()
