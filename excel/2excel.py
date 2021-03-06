#!/usr/bin/python3

import sys
from openpyxl import Workbook
import datetime

from openpyxl.chart import (LineChart, Reference,)
from openpyxl.chart.axis import DateAxis

wb = Workbook()
ws = wb.active
ws.title = "Memory Statistics"
file=[]
file.append(open('timer.memo','r'))
file.append(open('MF.VmData.memo', 'r'))
file.append(open('SMF.VmData.memo', 'r'))
file.append(open('MF.permemo.memo', 'r'))
file.append(open('SMF.permemo.memo', 'r'))
file.append(open('MF.percpu.memo', 'r'))
file.append(open('SMF.percpu.memo', 'r'))
file.append(open('MF.VmRss.memo', 'r'))
file.append(open('MF.heap.memo', 'r'))
file.append(open('MF.Thread.memo', 'r'))

title = ["Time", "MF VmData", "SMF VmData", "MF Memory %",  "SMF Memory %", "MF CPU %", "SMF CPU %", "MF VmRSS", "MF Heap", "MF Threads"]
data=[]
data.append(title)

for tValue in file[0]:
    date = datetime.datetime.strptime(tValue.strip(),'%b %d %H:%M:%S %Y')
    tmpline=[]
    tmpline.append(date)
    for col in range(1,10):
        value = str(file[col].readline()).strip()
        print ("col:%d, value:%s",col,value)
        if col < 3 and col >=7:
            value = int(value)
        else:
            value = float(value)
        tmpline.append(value)
    data.append(tmpline)
for row in data:
    ws.append(row)

data1 = Reference(ws,min_col=2, min_row=1, max_col=3, max_row=len(data))

c1 = LineChart()
c1.title = "Memory Information"
c1.style = 13
c1.y_axis.majorGridlines = None
c1.y_axis.title = "Memory (KB)"
c1.x_axis.title = "Time"

c1.add_data(data1, titles_from_data=True)

s1 = c1.series[0]
s1.graphicalProperties.line.solidFill = "FF0000"
#s1.graphicalProperties.line.dashStyle = "sysDot"
#s1.graphicalProperties.line.width = 100050 # width in EMUs

s2 = c1.series[1]
s2.graphicalProperties.line.solidFill = "990000"
s2.smooth = True

c2 = LineChart()
data2 = Reference(ws, min_col=4, min_row=1, max_col=7, max_row=len(data))
# c2.grouping = "percentStacked"
c2.add_data(data2, titles_from_data=True)
c2.y_axis.axId =200
c2.x_axis.title = "Time"
c2.y_axis.title = "CPU and Memory Ratio"
s21 = c2.series[0]
s21.graphicalProperties.line.solidFill = "0000FF"
#s21.graphicalProperties.line.dashStyle = "sysDash"
s21.smooth=True
s21.graphicalProperties.line.width = 30000 # width in EMUs
s22 = c2.series[1]
s22.graphicalProperties.line.solidFill = "000099"
#s22.graphicalProperties.line.dashStyle = "sysDash"
s22.graphicalProperties.line.width = 30000 # width in EMUs
s23 = c2.series[2]
s23.graphicalProperties.line.solidFill = "00FF00"
#s23.graphicalProperties.line.dashStyle = "sysDash"
s23.graphicalProperties.line.width = 30000 # width in EMUs
s24 = c2.series[3]
s24.graphicalProperties.line.solidFill = "009900"
#s24.graphicalProperties.line.dashStyle = "sysDash"
s24.graphicalProperties.line.width = 30000 # width in EMUs

c2.y_axis.crosses = "max"
c1 += c2
ws.add_chart(c1, "A"+str(len(data)+5))

data1x = Reference(ws,min_col=8, min_row=1, max_col=9, max_row=len(data))

c1x = LineChart()
c1x.title = "RSS Heap and Thread Information"
c1x.style = 13
c1x.y_axis.majorGridlines = None
c1x.y_axis.title = "Memory (KB)"
c1x.x_axis.title = "Time"

c1x.add_data(data1x, titles_from_data=True)

s1x = c1x.series[0]
s1x.graphicalProperties.line.solidFill = "FF0000"

s2x = c1x.series[1]
s2x.graphicalProperties.line.solidFill = "990000"
s2x.smooth = True

c2x = LineChart()
data2x = Reference(ws, min_col=10, min_row=1, max_col=10, max_row=len(data))
c2x.add_data(data2x, titles_from_data=True)
c2x.y_axis.axId =300
c2x.x_axis.title = "Time"
c2x.y_axis.title = "Thread"
s21x = c2x.series[0]
s21x.graphicalProperties.line.solidFill = "0000FF"
#s21x.graphicalProperties.line.dashStyle = "sysDash"
s21x.graphicalProperties.line.width = 30000 # width in EMUs

c2x.y_axis.crosses = "max"
c1x += c2x
ws.add_chart(c1x, "A"+str(len(data)+25))

wb.save(sys.argv[1]+".telemetry.xlsx")
for i in range(0,10):
    file[i].close()
